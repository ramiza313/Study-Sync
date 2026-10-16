
from __future__ import annotations
import json, os, re, uuid
from pathlib import Path
from datetime import datetime, timezone
from flask import Flask, request, jsonify, send_file
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

BASE = Path(__file__).resolve().parent
load_dotenv(BASE / ".env")
UPLOAD_DIR = BASE / "uploads"
DATA_DIR = BASE / "data"
UPLOAD_DIR.mkdir(exist_ok=True)
DATA_DIR.mkdir(exist_ok=True)
RESOURCE_DB = DATA_DIR / "resources.json"
MAX_BYTES = int(os.getenv("MAX_UPLOAD_MB", "20")) * 1024 * 1024
ALLOWED = {"pdf", "docx", "pptx", "txt", "md", "csv", "png", "jpg", "jpeg", "webp"}
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_BYTES


def load_db():
    if RESOURCE_DB.exists():
        try:
            return json.loads(RESOURCE_DB.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def save_db(data):
    RESOURCE_DB.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )


def extract_text(path: Path, ext: str) -> str:
    if ext in {"txt", "md", "csv"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == "pdf":
        import fitz
        with fitz.open(path) as doc:
            return "\n\n".join(page.get_text("text") for page in doc)
    if ext == "docx":
        from docx import Document
        doc = Document(path)
        bits = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                bits.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(bits)
    if ext == "pptx":
        from pptx import Presentation
        prs = Presentation(path)
        bits = []
        for i, slide in enumerate(prs.slides, 1):
            parts = [
                shape.text.strip()
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
                and shape.text.strip()
            ]
            if parts:
                bits.append(f"Slide {i}: " + "\n".join(parts))
        return "\n\n".join(bits)
    if ext in {"png", "jpg", "jpeg", "webp"}:
        try:
            from PIL import Image
            import pytesseract
            return pytesseract.image_to_string(Image.open(path))
        except Exception as exc:
            raise ValueError(
                "Image OCR needs Pillow, pytesseract, and the Tesseract OCR "
                "engine installed. Details: " + str(exc)
            )
    raise ValueError("Unsupported file type")


@app.get("/")
def home():
    return send_file(BASE / "index.html")


@app.get("/api/health")
def health():
    return jsonify({
        "ok": True,
        "ai_configured": bool(os.getenv("GEMMA_API_KEY")),
        "model": os.getenv("GEMMA_MODEL", "gemma-4"),
        "message": "Backend is running. Add GEMMA_API_KEY to enable AI generation."
    })


@app.post("/api/resources")
def upload_resource():
    file = request.files.get("file")
    subject_id = re.sub(
        r"[^a-zA-Z0-9_-]", "", request.form.get("subject_id", "")
    )
    subject_name = request.form.get("subject_name", "")[:160]

    if not file or not file.filename:
        return jsonify(error="Choose a file to upload."), 400
    if not subject_id:
        return jsonify(error="A subject ID is required."), 400

    filename = secure_filename(file.filename)
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext not in ALLOWED:
        return jsonify(
            error=f"Unsupported file type: .{ext}. Supported types: "
                  f"{', '.join(sorted(ALLOWED))}."
        ), 415

    resource_id = uuid.uuid4().hex
    subject_dir = UPLOAD_DIR / subject_id
    subject_dir.mkdir(parents=True, exist_ok=True)
    saved_path = subject_dir / f"{resource_id}_{filename}"
    file.save(saved_path)

    try:
        extracted = extract_text(saved_path, ext).strip()
    except Exception as exc:
        saved_path.unlink(missing_ok=True)
        return jsonify(error=str(exc)), 422

    if len(extracted) < 20:
        saved_path.unlink(missing_ok=True)
        return jsonify(
            error="Very little text could be extracted. For scanned PDFs, "
                  "use OCR or upload a clearer image/document."
        ), 422

    db = load_db()
    db.setdefault(subject_id, []).append({
        "resource_id": resource_id,
        "subject_id": subject_id,
        "subject_name": subject_name,
        "filename": filename,
        "stored_path": str(saved_path.relative_to(BASE)),
        "extension": ext,
        "text": extracted[:250000],
        "chars": len(extracted),
        "uploaded_at": datetime.now(timezone.utc).isoformat()
    })
    save_db(db)

    return jsonify(
        resource_id=resource_id,
        filename=filename,
        chars=len(extracted),
        status="Text extracted and saved",
        subject_id=subject_id
    )


# ADDED: Delete a teaching resource and its extracted text.
@app.delete("/api/resources/<resource_id>")
def delete_resource(resource_id):
    db = load_db()

    # Find the resource and the subject it belongs to.
    found_subject_id = None
    found_resource = None

    for subject_id, resources in db.items():
        for resource in resources:
            if resource.get("resource_id") == resource_id:
                found_subject_id = subject_id
                found_resource = resource
                break
        if found_resource:
            break

    if not found_resource:
        return jsonify(error="Resource not found."), 404

    # Only delete files located inside the configured uploads directory.
    stored_path = found_resource.get("stored_path")
    if stored_path:
        try:
            file_path = (BASE / stored_path).resolve()
            upload_root = UPLOAD_DIR.resolve()

            if file_path != upload_root and upload_root in file_path.parents:
                file_path.unlink(missing_ok=True)
            else:
                return jsonify(
                    error="The resource file path is invalid; nothing was deleted."
                ), 400
        except (OSError, ValueError) as exc:
            return jsonify(error=f"Could not delete the uploaded file: {exc}"), 500

    # Remove the resource record and its extracted text from resources.json.
    db[found_subject_id] = [
        resource for resource in db[found_subject_id]
        if resource.get("resource_id") != resource_id
    ]

    if not db[found_subject_id]:
        del db[found_subject_id]

    save_db(db)

    # Clear saved drafts for this subject so new packs can be generated
    # from the remaining teaching resources.
    packs_path = DATA_DIR / "learning_packs.json"
    if packs_path.exists():
        try:
            packs = json.loads(packs_path.read_text(encoding="utf-8"))
            packs = {
                key: value
                for key, value in packs.items()
                if not (
                    isinstance(value, dict)
                    and value.get("subject_id") == found_subject_id
                )
            }
            packs_path.write_text(
                json.dumps(packs, ensure_ascii=False, indent=2),
                encoding="utf-8"
            )
        except (OSError, json.JSONDecodeError) as exc:
            app.logger.warning(
                "Resource deleted, but learning-pack drafts could not be "
                "cleaned up: %s", exc
            )

    return jsonify(
        ok=True,
        resource_id=resource_id,
        subject_id=found_subject_id,
        message="Resource and extracted text deleted successfully."
    )


def call_gemma(prompt: str) -> str:
    api_key = os.getenv("GEMMA_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if not api_key:
        raise RuntimeError(
            "GEMMA_API_KEY is missing. Copy .env.example to .env "
            "and add your API key."
        )
    from google import genai
    client = genai.Client(api_key=api_key)
    model = os.getenv("GEMMA_MODEL", "gemma-4")
    response = client.models.generate_content(model=model, contents=prompt)
    return (getattr(response, "text", None) or "").strip()


def parse_json_response(raw: str) -> dict:
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip(), flags=re.I)
    match = re.search(r"\{.*\}", cleaned, flags=re.S)
    if not match:
        raise ValueError("The model did not return valid JSON. Please try again.")
    obj = json.loads(match.group(0))
    obj.setdefault("summary", "")
    obj.setdefault("activity", "")
    obj.setdefault("quiz", [])
    obj.setdefault("sources", [])
    return obj


@app.post("/api/generate")
def generate_pack():
    payload = request.get_json(silent=True) or {}
    subject_id = re.sub(
        r"[^a-zA-Z0-9_-]", "", str(payload.get("subject_id", ""))
    )
    subject_name = str(payload.get("subject_name", "Selected subject"))[:160]
    topic_name = str(payload.get("topic_name", "Selected topic"))[:200]

    if not subject_id or not topic_name:
        return jsonify(error="Subject and topic are required."), 400

    db = load_db()
    resources = db.get(subject_id, [])
    if not resources:
        return jsonify(
            error="Upload at least one teaching source for this subject "
                  "before generating a pack."
        ), 400

    # Include bounded excerpts, labelled by original filename,
    # so output can be checked against sources.
    source_blocks = []
    for r in resources:
        excerpt = r.get("text", "")[:12000]
        if excerpt:
            source_blocks.append(f"SOURCE FILE: {r['filename']}\n{excerpt}")

    context = "\n\n--- SOURCE SEPARATOR ---\n\n".join(source_blocks)[:45000]

    prompt = f'''You are an educational assistant creating a learning pack for a college student who may have missed a lecture. Use ONLY the provided teacher-uploaded source excerpts for factual claims. If the sources do not contain enough information, say so in the summary and do not invent facts. Keep it clear and suitable for a first-year undergraduate student.

Subject: {subject_name}
Selected syllabus topic: {topic_name}

Teacher's source excerpts:
{context}

Return ONLY valid JSON, no markdown, with this schema:
{{"summary":"A clear topic-specific explanation in 2-5 short paragraphs","activity":"One practical learning activity with instructions and expected output","quiz":[{{"question":"...","options":["A. ...","B. ...","C. ...","D. ..."],"answer":"exact option text","explanation":"why the answer is correct, based on the source"}}],"sources":["filename — brief note about relevant passage"]}}
Create 3 multiple-choice questions. Ensure each answer exactly matches one option. Do not include unsupported claims. Use the selected topic, not a generic subject overview.'''

    try:
        raw = call_gemma(prompt)
        pack = parse_json_response(raw)

        # Store the generated draft server-side as well,
        # without marking it approved.
        pack_key = f"{subject_id}::{payload.get('topic_id','')}"
        packs_path = DATA_DIR / "learning_packs.json"
        packs = (
            json.loads(packs_path.read_text(encoding="utf-8"))
            if packs_path.exists() else {}
        )
        packs[pack_key] = {
            "subject_id": subject_id,
            "topic_id": payload.get("topic_id", ""),
            "topic_name": topic_name,
            "pack": pack,
            "approved": False,
            "generated_at": datetime.now(timezone.utc).isoformat()
        }
        packs_path.write_text(
            json.dumps(packs, ensure_ascii=False, indent=2),
            encoding="utf-8"
        )

        return jsonify(
            pack=pack,
            topic_name=topic_name,
            sources_used=[r["filename"] for r in resources],
            model=os.getenv("GEMMA_MODEL", "gemma-4"),
            approved=False
        )
    except Exception as exc:
        return jsonify(error=str(exc)), 502


@app.errorhandler(413)
def too_large(_):
    return jsonify(
        error=f"File too large. Maximum upload size is "
              f"{MAX_BYTES // (1024*1024)} MB."
    ), 413


if __name__ == "__main__":
    app.run(
        host="127.0.0.1",
        port=int(os.getenv("PORT", "5000")),
        debug=True
    )
